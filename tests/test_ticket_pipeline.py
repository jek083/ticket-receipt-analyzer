import csv
import tempfile
import unittest
from pathlib import Path
from unittest.mock import Mock, patch

import config
from openpyxl import load_workbook
from pypdf.errors import PdfReadError
from src.analyzer import TicketAnalyzer
from src.parse_ticket_enrichi import categorize_item, parse_ticket
from src.parser_dispatcher import parse_ticket as dispatch_ticket
from src.pdf_extractor import PDFExtractor
from src.ticket_io import export_items_csv, load_tickets_json, write_tickets_json


OCR_TEXT = """E. Leclerc
28 juin 2024 10:12
--- Page 1 (OCR) ---
>> EPICERIE
LAIT TGU 1/2 ECREME 1L 1,01 1
PILSNER (6X33cl)
6 X 1,19 € 7,14 3
>> FRUITS ET LEGUMES
COURGETTE LONGUE VERTE
0,828 kg X 1,76 €/kg
1,46 1
Total 3 artirlec Qf 9
Total 8 articles 9,61
"""

PICARD_OCR_TEXT = """1067 - PICARD SURGELES SAS
V1
V1
V1
V1 *Gl Amarena It 2,99 €
V1 *Epinards bio 2,90 €
V1 *Naanwich poulet 3,99 €
Sélection Picard & Nous -0,80 €
V1 *Chapati sce méch 3,99 €
--------------
Total sans remise
13,87 €
Total des remises
-0,80 €
--------------
TOTAL (4) 13,07 €
Date
09.05.26 17:17
Nb lignes ticket: 4
"""


class EnrichedParserTests(unittest.TestCase):
    def test_picard_ice_cream_keywords_use_glaces_category(self):
        self.assertEqual(
            categorize_item("Gl Amarena It", config.FOOD_CATEGORIES),
            "glaces",
        )
        self.assertEqual(
            categorize_item("Nocciola crème g", config.FOOD_CATEGORIES),
            "glaces",
        )

    def test_dispatcher_keeps_leclerc_parser(self):
        ticket = dispatch_ticket(OCR_TEXT, source_file="leclerc.pdf")

        self.assertEqual(ticket["fichier_source"], "leclerc.pdf")
        self.assertEqual(ticket["articles_detectes"], 3)
        self.assertEqual(ticket["articles"][0]["article"], "LAIT TGU 1/2 ECREME 1L")

    def test_dispatcher_recognizes_auredis_as_leclerc(self):
        text = OCR_TEXT.replace("E. Leclerc", "5.A.S AUREDIS")
        ticket = dispatch_ticket(text, source_file="auredis.pdf")

        self.assertEqual(ticket["articles_detectes"], 3)

    def test_parses_ticket_metadata_and_one_and_multiline_items(self):
        ticket = parse_ticket(
            OCR_TEXT,
            source_file="leclerc.pdf",
            categories={
                "produits_laitiers": ["lait"],
                "boissons": ["bière", "pilsner"],
            },
        )

        self.assertEqual(ticket["fichier_source"], "leclerc.pdf")
        self.assertEqual(ticket["enseigne"], "E.Leclerc")
        self.assertEqual(ticket["date"], "2024-06-28 10:12")
        self.assertEqual(ticket["total_ticket"], 9.61)
        self.assertEqual(ticket["nombre_articles_ticket"], 8)
        self.assertEqual(ticket["articles_detectes"], 3)
        self.assertEqual(ticket["unites_detectees"], 8)

        milk, beer, courgette = ticket["articles"]
        self.assertEqual(milk["article"], "LAIT TGU 1/2 ECREME 1L")
        self.assertEqual(milk["prix_total"], 1.01)
        self.assertEqual(milk["categorie"], "produits_laitiers")
        self.assertEqual(beer["quantite"], 6)
        self.assertEqual(beer["prix_unitaire"], 1.19)
        self.assertEqual(beer["prix_total"], 7.14)
        self.assertEqual(beer["tva_code"], "3")
        self.assertEqual(courgette["quantite"], 0.828)
        self.assertEqual(courgette["poids_kg"], 0.828)
        self.assertEqual(courgette["prix_kg"], 1.76)
        self.assertEqual(courgette["rayon"], "FRUITS ET LEGUMES")

    def test_parses_compact_product_line(self):
        ticket = parse_ticket(
            "Magasin\n>> CREMERIE\nYAOURT NATURE 0,89 1\nTotal 1 article 0,89"
        )

        self.assertEqual(ticket["articles"][0]["article"], "YAOURT NATURE")
        self.assertEqual(ticket["articles"][0]["prix_total"], 0.89)
        self.assertEqual(ticket["total_ticket"], 0.89)

    def test_parses_numeric_date_formats(self):
        for date_text in ("2024-06-28", "28-06-2024"):
            with self.subTest(date_text=date_text):
                ticket = parse_ticket(
                    "Magasin\n{}\n>> CREMERIE\nYAOURT 0,89 1".format(date_text)
                )
                self.assertEqual(ticket["date"], "2024-06-28")

    def test_rejects_empty_or_unparseable_ticket_text(self):
        with self.assertRaisesRegex(ValueError, "vide"):
            parse_ticket(" \n--- Page 1 (OCR) ---\n")
        with self.assertRaisesRegex(ValueError, "Aucun article"):
            parse_ticket("E. Leclerc\n28/06/2024\nAucun détail")


class PicardParserTests(unittest.TestCase):
    def test_parses_picard_ticket_into_shared_schema(self):
        ticket = dispatch_ticket(
            PICARD_OCR_TEXT,
            source_file="picard.pdf",
            categories={
                "fruits_legumes": ["epinards"],
                "viandes": ["poulet"],
            },
        )

        self.assertEqual(ticket["fichier_source"], "picard.pdf")
        self.assertEqual(ticket["enseigne"], "Picard")
        self.assertEqual(ticket["date"], "2026-05-09 17:17")
        self.assertEqual(ticket["total_ticket"], 13.07)
        self.assertEqual(ticket["total_remises"], 0.80)
        self.assertEqual(ticket["nombre_articles_ticket"], 4)
        self.assertEqual(ticket["articles_detectes"], 4)
        self.assertEqual(ticket["unites_detectees"], 4)

        items = ticket["articles"]
        self.assertEqual(
            [item["article"] for item in items],
            [
                "Gl Amarena It",
                "Epinards bio",
                "Naanwich poulet",
                "Chapati sce méch",
            ],
        )
        self.assertAlmostEqual(
            sum(item["prix_total"] for item in items),
            13.87,
        )
        self.assertAlmostEqual(
            sum(item["prix_total"] for item in items) - ticket["total_remises"],
            ticket["total_ticket"],
        )
        self.assertTrue(all(item["rayon"] == "SURGELES" for item in items))
        self.assertEqual(items[1]["categorie"], "fruits_legumes")
        self.assertEqual(items[2]["categorie"], "viandes")
        self.assertIsNone(items[0]["categorie"])
        self.assertTrue(
            all(
                item["poids_kg"] is None
                and item["prix_kg"] is None
                and item["tva_code"] is None
                for item in items
            )
        )

    def test_parses_picard_name_and_price_on_separate_lines(self):
        text = PICARD_OCR_TEXT.replace(
            "V1 *Gl Amarena It 2,99 €",
            "V1 *Gl Amarena It\n2,99 €",
        )
        ticket = dispatch_ticket(text)

        self.assertEqual(ticket["articles"][0]["article"], "Gl Amarena It")
        self.assertEqual(ticket["articles"][0]["prix_total"], 2.99)
        self.assertEqual(ticket["articles_detectes"], 4)


class TicketIOTests(unittest.TestCase):
    def test_json_is_reloaded_before_flat_csv_export(self):
        ticket = parse_ticket(OCR_TEXT, source_file="leclerc.pdf")
        with tempfile.TemporaryDirectory() as temp_dir:
            json_path = Path(temp_dir) / "processed" / "tickets.json"
            csv_path = Path(temp_dir) / "output" / "items.csv"

            write_tickets_json([ticket], json_path)
            reloaded = load_tickets_json(json_path)
            export_items_csv(reloaded, csv_path)

            with csv_path.open(encoding="utf-8-sig", newline="") as source:
                rows = list(csv.DictReader(source, delimiter=";"))

        self.assertEqual(len(rows), 3)
        self.assertEqual(rows[1]["article"], "PILSNER (6X33cl)")
        self.assertEqual(rows[1]["date"], "2024-06-28 10:12")
        self.assertEqual(rows[1]["enseigne"], "E.Leclerc")
        self.assertEqual(rows[1]["total_remises"], "")
        self.assertEqual(rows[1]["quantite"], "6")
        self.assertEqual(rows[1]["unites_detectees"], "8")
        self.assertEqual(rows[2]["rayon"], "FRUITS ET LEGUMES")
        self.assertEqual(rows[2]["poids_kg"], "0.828")

    def test_picard_json_keeps_unknown_fields_as_null(self):
        ticket = dispatch_ticket(PICARD_OCR_TEXT, source_file="picard.pdf")
        with tempfile.TemporaryDirectory() as temp_dir:
            json_path = Path(temp_dir) / "processed" / "tickets.json"
            write_tickets_json([ticket], json_path)
            reloaded = load_tickets_json(json_path)

        self.assertEqual(reloaded[0]["total_remises"], 0.80)
        self.assertIsNone(reloaded[0]["articles"][0]["categorie"])
        self.assertIsNone(reloaded[0]["articles"][0]["tva_code"])

class AnalyzerTests(unittest.TestCase):
    def test_includes_article_rows_from_every_ticket(self):
        tickets = [
            {
                "fichier_source": "first.pdf",
                "enseigne": "Leclerc",
                "date": "2024-06-28",
                "total_ticket": 2.0,
                "articles": [
                    {
                        "article": "LAIT",
                        "prix_total": 2.0,
                        "categorie": "produits_laitiers",
                    }
                ],
            },
            {
                "fichier_source": "second.pdf",
                "enseigne": "Auchan",
                "date": "2024-06-29",
                "total_ticket": 3.0,
                "articles": [
                    {
                        "article": "PAIN",
                        "prix_total": 3.0,
                        "categorie": "boulangerie",
                    }
                ],
            },
        ]

        analyzer = TicketAnalyzer(tickets)

        self.assertEqual(
            set(analyzer.df["item_name"].dropna()),
            {"LAIT", "PAIN"},
        )
        self.assertEqual(len(analyzer.summary_by_store()), 2)

    def test_excel_export_uses_french_labels_and_categories(self):
        tickets = [
            {
                "fichier_source": "ticket.pdf",
                "enseigne": "Picard",
                "date": "2026-05-09",
                "total_ticket": 5.98,
                "articles": [
                    {
                        "article": "Nocciola crème",
                        "prix_total": 2.99,
                        "categorie": "produits_laitiers",
                    },
                    {
                        "article": "Gl Amarena",
                        "prix_total": 2.99,
                        "categorie": "glaces",
                    },
                ],
            }
        ]

        with tempfile.TemporaryDirectory() as temp_dir:
            excel_path = Path(temp_dir) / "analyse_tickets.xlsx"
            TicketAnalyzer(tickets).export_to_excel(str(excel_path))

            workbook = load_workbook(excel_path, read_only=True, data_only=True)
            try:
                self.assertEqual(
                    workbook.sheetnames,
                    [
                        "Par enseigne",
                        "Par catégorie",
                        "Tendance temporelle",
                        "Évolution par enseigne",
                        "Prix par catégorie et enseigne",
                        "Top 20 des articles",
                    ],
                )
                category_values = {
                    cell.value
                    for row in workbook["Par catégorie"].iter_rows()
                    for cell in row
                }
                price_values = {
                    cell.value
                    for row in workbook["Prix par catégorie et enseigne"].iter_rows()
                    for cell in row
                }
                store_values = {
                    cell.value
                    for row in workbook["Par enseigne"].iter_rows()
                    for cell in row
                }
                top_item_values = {
                    cell.value
                    for row in workbook["Top 20 des articles"].iter_rows()
                    for cell in row
                }
            finally:
                workbook.close()

        self.assertIn("Catégorie", category_values)
        self.assertIn("Produits laitiers", category_values)
        self.assertIn("Exemples d’articles", category_values)
        self.assertIn("Glaces", price_values)
        self.assertIn("Enseigne", store_values)
        self.assertIn("Dépenses par ticket (€)", store_values)
        self.assertIn("Nombre d’achats", top_item_values)


class PDFExtractorTests(unittest.TestCase):
    def test_continues_the_batch_and_records_failed_pdfs(self):
        with tempfile.TemporaryDirectory() as temp_dir:
            folder = Path(temp_dir)
            (folder / "good.pdf").touch()
            (folder / "bad.pdf").touch()
            extractor = PDFExtractor()

            def extract(path):
                if path.name == "bad.pdf":
                    raise ValueError("invalid PDF")
                return "Recognized ticket"

            with patch.object(extractor, "extract_from_file", side_effect=extract):
                results = extractor.extract_from_folder(folder)

        self.assertEqual(results, {"good.pdf": "Recognized ticket"})
        self.assertEqual(extractor.failed_files, {"bad.pdf": "invalid PDF"})

    def test_renders_ocr_pages_at_configured_resolution(self):
        page = Mock()
        page.extract_text.return_value = None
        page.to_image.return_value.original = object()
        pypdf_page = Mock()
        pypdf_page.extract_text.return_value = ""
        pypdf_reader = Mock(pages=[pypdf_page])
        pdf = Mock()
        pdf.pages = [page]
        pdf_context = Mock()
        pdf_context.__enter__ = Mock(return_value=pdf)
        pdf_context.__exit__ = Mock(return_value=False)

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "ticket.pdf"
            pdf_path.touch()
            with patch("src.pdf_extractor.pdfplumber.open", return_value=pdf_context):
                with patch("src.pdf_extractor.PdfReader", return_value=pypdf_reader):
                    with patch(
                        "src.pdf_extractor.pytesseract.image_to_string",
                        return_value="Ticket OCR",
                    ):
                        text = PDFExtractor(
                            ocr_resolution=250
                        ).extract_from_file(pdf_path)

        self.assertEqual(text, "Ticket OCR")
        page.to_image.assert_called_once_with(resolution=250)

    def test_skips_an_empty_page_if_other_pages_have_text(self):
        pages = []
        for _ in range(2):
            page = Mock()
            page.extract_text.return_value = None
            page.to_image.return_value.original = object()
            pages.append(page)

        pypdf_pages = [Mock(), Mock()]
        for pypdf_page in pypdf_pages:
            pypdf_page.extract_text.return_value = ""
        pypdf_reader = Mock(pages=pypdf_pages)

        pdf = Mock()
        pdf.pages = pages
        pdf_context = Mock()
        pdf_context.__enter__ = Mock(return_value=pdf)
        pdf_context.__exit__ = Mock(return_value=False)

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "ticket.pdf"
            pdf_path.touch()
            with patch("src.pdf_extractor.pdfplumber.open", return_value=pdf_context):
                with patch("src.pdf_extractor.PdfReader", return_value=pypdf_reader):
                    with patch(
                        "src.pdf_extractor.pytesseract.image_to_string",
                        side_effect=["Ticket OCR", ""],
                    ):
                        text = PDFExtractor().extract_from_file(pdf_path)

        self.assertEqual(text, "Ticket OCR")

    def test_fails_if_all_pages_are_empty(self):
        page = Mock()
        page.extract_text.return_value = None
        page.to_image.return_value.original = object()
        pypdf_page = Mock()
        pypdf_page.extract_text.return_value = ""
        pypdf_reader = Mock(pages=[pypdf_page])
        pdf = Mock()
        pdf.pages = [page]
        pdf_context = Mock()
        pdf_context.__enter__ = Mock(return_value=pdf)
        pdf_context.__exit__ = Mock(return_value=False)

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "ticket.pdf"
            pdf_path.touch()
            with patch("src.pdf_extractor.pdfplumber.open", return_value=pdf_context):
                with patch("src.pdf_extractor.PdfReader", return_value=pypdf_reader):
                    with patch(
                        "src.pdf_extractor.pytesseract.image_to_string",
                        return_value="",
                    ):
                        with self.assertRaisesRegex(
                            ValueError,
                            "Aucun texte extrait",
                        ):
                            PDFExtractor().extract_from_file(pdf_path)

    def test_uses_pypdf_layout_text_before_ocr(self):
        page = Mock()
        page.extract_text.return_value = None
        pypdf_page = Mock()
        pypdf_page.extract_text.return_value = "Native ticket text"
        pypdf_reader = Mock(pages=[pypdf_page])
        pdf = Mock(pages=[page])
        pdf_context = Mock()
        pdf_context.__enter__ = Mock(return_value=pdf)
        pdf_context.__exit__ = Mock(return_value=False)

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "ticket.pdf"
            pdf_path.touch()
            with patch("src.pdf_extractor.pdfplumber.open", return_value=pdf_context):
                with patch(
                    "src.pdf_extractor.PdfReader",
                    return_value=pypdf_reader,
                ) as reader_factory:
                    with patch(
                        "src.pdf_extractor.pytesseract.image_to_string"
                    ) as ocr:
                        text = PDFExtractor(ocr_enabled=False).extract_from_file(
                            pdf_path
                        )

        self.assertEqual(text, "Native ticket text")
        reader_factory.assert_called_once_with(str(pdf_path), strict=False)
        pypdf_page.extract_text.assert_called_once_with(
            extraction_mode="layout"
        )
        ocr.assert_not_called()

    def test_uses_ocr_if_pypdf_cannot_read_the_pdf(self):
        page = Mock()
        page.extract_text.return_value = None
        page.to_image.return_value.original = object()
        pdf = Mock(pages=[page])
        pdf_context = Mock()
        pdf_context.__enter__ = Mock(return_value=pdf)
        pdf_context.__exit__ = Mock(return_value=False)

        with tempfile.TemporaryDirectory() as temp_dir:
            pdf_path = Path(temp_dir) / "ticket.pdf"
            pdf_path.touch()
            with patch("src.pdf_extractor.pdfplumber.open", return_value=pdf_context):
                with patch(
                    "src.pdf_extractor.PdfReader",
                    side_effect=PdfReadError("invalid PDF"),
                ):
                    with patch(
                        "src.pdf_extractor.pytesseract.image_to_string",
                        return_value="Ticket OCR",
                    ):
                        text = PDFExtractor().extract_from_file(pdf_path)

        self.assertEqual(text, "Ticket OCR")
